import os
import shutil
import docx
from docx.shared import Inches, Pt, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DESKTOP_DOCX = r"C:\Users\itzmo\Desktop\ผลิตสื่อ.docx"
BACKUP_DOCX = r"C:\Users\itzmo\Desktop\ผลิตสื่อ.backup.docx"
FALLBACK_DOCX = r"C:\Users\itzmo\Desktop\ผลิตสื่อ_อัปเดตหน้าจอเว็บ.docx"
REPO_DOCX = r"d:\VsCode\Project\Content-Management-System\docs\academic\reports\SRS-ผลิตสื่อ-สมบูรณ์.docx"
SCREENSHOT_DIR = r"d:\VsCode\Project\Content-Management-System\docs\screenshots\admin-web"

ADMIN_SCREENS = [
    {
        "sub_no": "3.2.1",
        "title": "หน้าเข้าสู่ระบบ",
        "desc": "หน้าที่ผู้ดูแลระบบทำการเข้าสู่ระบบเว็บแอปพลิเคชัน โดยกรอกอีเมลผู้ดูแลระบบและรหัสผ่าน จากนั้นกดปุ่มเข้าสู่ระบบแดชบอร์ด หากกรอกข้อมูลไม่ถูกต้องระบบจะแจ้งเตือนให้ตรวจสอบข้อมูลอีกครั้ง",
        "img": "01_admin_login.png",
        "fig_no": 46,
    },
    {
        "sub_no": "3.2.2",
        "title": "หน้าแดชบอร์ดภาพรวมระบบ",
        "desc": "หน้าที่ผู้ดูแลระบบสามารถดูข้อมูลสรุปภาพรวมของระบบทั้งหมด เช่น จำนวนผู้ใช้งานในระบบ จำนวนทีมผลิตสื่อ ชิ้นงานทั้งหมด งานย่อยที่กำลังดำเนินการ สถานะความพร้อมของเซิร์ฟเวอร์และฐานข้อมูล รวมถึงรายการกิจกรรมล่าสุดที่เกิดขึ้นในระบบ",
        "img": "02_admin_dashboard.png",
        "fig_no": 47,
    },
    {
        "sub_no": "3.2.3",
        "title": "หน้าจัดการข้อมูลผู้ใช้งานและสิทธิ์ระบบ",
        "desc": "หน้าที่ผู้ดูแลระบบใช้ตรวจสอบรายชื่อบัญชีผู้ใช้งานทั้งหมดในระบบ สามารถค้นหาชื่อหรืออีเมล กรองตามบทบาทหรือทีม กำหนดสิทธิ์การใช้งาน (ผู้ดูแลระบบ ผู้จัดการทีม หรือสมาชิก) เปลี่ยนทีมที่สังกัด รวมถึงระงับการใช้งานหรือลบบัญชีผู้ใช้ได้",
        "img": "03_admin_users_list.png",
        "fig_no": 48,
    },
    {
        "sub_no": "3.2.4",
        "title": "หน้าเพิ่มผู้ใช้งานใหม่",
        "desc": "หน้าที่ผู้ดูแลระบบสามารถสร้างบัญชีผู้ใช้งานใหม่เข้าสู่ระบบโดยตรง โดยกรอกชื่อ-นามสกุล อีเมล รหัสผ่านเริ่มต้น กำหนดบทบาทหน้าที่ในระบบ และเลือกทีมที่ต้องการให้สังกัด จากนั้นกดปุ่มสร้างบัญชีผู้ใช้",
        "img": "04_admin_add_user_modal.png",
        "fig_no": 49,
    },
    {
        "sub_no": "3.2.5",
        "title": "หน้าจัดการโครงสร้างทีม",
        "desc": "หน้าที่ผู้ดูแลระบบสามารถดูรายการทีมผลิตสื่อทั้งหมดในระบบ ตรวจสอบรหัสเข้าร่วมทีม (Invite Code) ดูรายชื่อผู้จัดการทีมและจำนวนสมาชิกในแต่ละทีม รวมถึงสามารถแก้ไขข้อมูลทีมหรือลบทีมออกจากระบบได้",
        "img": "05_admin_teams_list.png",
        "fig_no": 50,
    },
    {
        "sub_no": "3.2.6",
        "title": "หน้าสร้างโครงสร้างทีมใหม่",
        "desc": "หน้าที่ผู้ดูแลระบบใช้สำหรับสร้างทีมผลิตสื่อใหม่ โดยระบุชื่อทีม รหัสเข้าร่วมทีมสำหรับให้สมาชิกใช้กรอกในโมบายแอปพลิเคชัน คำอธิบายรายละเอียดของทีม และเลือกแต่งตั้งผู้จัดการทีมประจำกลุ่มงาน",
        "img": "06_admin_create_team_modal.png",
        "fig_no": 51,
    },
    {
        "sub_no": "3.2.7",
        "title": "หน้ากำกับดูแลชิ้นงานทั้งหมด",
        "desc": "หน้าที่ผู้ดูแลระบบสามารถตรวจสอบรายการชิ้นงานการผลิตสื่อทั้งหมดของทุกทีมในระบบ สามารถค้นหาชื่อชิ้นงาน กรองตามสถานะการผลิต ดูทีมที่รับผิดชอบ กำหนดการส่งมอบ ลิงก์ผลงาน รวมถึงแก้ไขสถานะหรือลบชิ้นงานได้",
        "img": "07_admin_contents_list.png",
        "fig_no": 52,
    },
    {
        "sub_no": "3.2.8",
        "title": "หน้าสร้างชิ้นงานใหม่",
        "desc": "หน้าที่ผู้ดูแลระบบสามารถสร้างรายการชิ้นงานใหม่จากส่วนกลาง โดยกรอกชื่อชิ้นงาน รายละเอียดขอบเขตงาน เลือกทีมผลิตสื่อที่รับผิดชอบ กำหนดสถานะเริ่มต้น และระบุวันกำหนดส่งมอบงาน",
        "img": "08_admin_create_content_modal.png",
        "fig_no": 53,
    },
    {
        "sub_no": "3.2.9",
        "title": "หน้าจัดการภารกิจและงานย่อยทั้งหมด",
        "desc": "หน้าที่ผู้ดูแลระบบใช้ติดตามงานย่อยทั้งหมดที่ถูกมอบหมายให้สมาชิกในแต่ละทีม สามารถค้นหาชื่องาน กรองตามสถานะการทำงาน (รอดำเนินการ กำลังทำ รอตรวจ หรือเสร็จสิ้น) ตรวจสอบผู้รับผิดชอบ ชิ้นงานหลักที่เชื่อมโยง และจัดการลบรายการงานย่อยได้",
        "img": "09_admin_tasks_list.png",
        "fig_no": 54,
    },
    {
        "sub_no": "3.2.10",
        "title": "หน้าจัดการข้อมูลประเภทงานย่อย",
        "desc": "หน้าที่ผู้ดูแลระบบใช้กำหนดและจัดการหมวดหมู่ประเภทงานย่อยในกระบวนการผลิตสื่อ เช่น งานตัดต่อวิดีโอ งานออกแบบกราฟิก งานเขียนบท หรือการถ่ายทำ สามารถเพิ่มประเภทงานใหม่ แก้ไขชื่อและคำอธิบาย หรือลบประเภทงานที่ไม่ใช้งานได้",
        "img": "10_admin_task_types.png",
        "fig_no": 55,
    },
    {
        "sub_no": "3.2.11",
        "title": "หน้าตรวจสอบประวัติการดำเนินงานของผู้ใช้งาน",
        "desc": "หน้าที่ผู้ดูแลระบบใช้ตรวจสอบบันทึกประวัติการใช้งานและการเปลี่ยนแปลงข้อมูลในระบบย้อนหลัง (Audit Logs) โดยแสดงวันเวลา ผู้ดำเนินการ ประเภทกิจกรรม และรายละเอียดการทำรายการ พร้อมทั้งสามารถค้นหาและกรองตามประเภทเหตุการณ์ได้",
        "img": "11_admin_audit_logs.png",
        "fig_no": 56,
    },
    {
        "sub_no": "3.2.12",
        "title": "หน้าตรวจสอบสถานะระบบและตั้งค่า",
        "desc": "หน้าที่ผู้ดูแลระบบใช้ตรวจสอบสถานะการทำงานของเซิร์ฟเวอร์ การเชื่อมต่อฐานข้อมูล MongoDB ข้อมูลสภาพแวดล้อมของระบบ เวลาการทำงานต่อเนื่องของเซิร์ฟเวอร์ และข้อมูลบัญชีผู้ดูแลระบบที่กำลังใช้งานอยู่",
        "img": "12_admin_settings_health.png",
        "fig_no": 57,
    },
]


def set_thai_font(run, size_pt=16, bold=False):
    run.font.name = "TH Sarabun New"
    run.font.size = Pt(size_pt)
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), "TH Sarabun New")
    rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
    rFonts.set(qn("w:cs"), "TH Sarabun New")

    sz_val = str(int(size_pt * 2))
    sz = rPr.find(qn("w:sz"))
    if sz is None:
        sz = OxmlElement("w:sz")
        rPr.append(sz)
    sz.set(qn("w:val"), sz_val)

    szCs = rPr.find(qn("w:szCs"))
    if szCs is None:
        szCs = OxmlElement("w:szCs")
        rPr.append(szCs)
    szCs.set(qn("w:val"), sz_val)

    if bold:
        bCs = rPr.find(qn("w:bCs"))
        if bCs is None:
            bCs = OxmlElement("w:bCs")
            rPr.append(bCs)


def insert_paragraph_before(ref_p, style_name="Normal"):
    new_p = OxmlElement("w:p")
    ref_p._p.addprevious(new_p)
    p = docx.text.paragraph.Paragraph(new_p, ref_p._parent)
    if style_name:
        p.style = ref_p.part.styles[style_name]
    return p


def add_seq_caption(p, fig_no, title_text):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run("ภาพที่ ")
    set_thai_font(r1, 16, bold=False)

    # SEQ field begin
    r_begin = p.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    r_begin._r.append(fldChar1)

    r_instr = p.add_run()
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " SEQ ภาพที่ \\* ARABIC "
    r_instr._r.append(instrText)

    r_sep = p.add_run()
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    r_sep._r.append(fldChar2)

    r_val = p.add_run(str(fig_no))
    set_thai_font(r_val, 16, bold=False)

    r_end = p.add_run()
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")
    r_end._r.append(fldChar3)

    r_space = p.add_run(" ")
    set_thai_font(r_space, 16, bold=False)

    r_title = p.add_run(title_text)
    set_thai_font(r_title, 16, bold=True)


def update_seq_number_in_paragraph(p, new_num):
    # Locate the run inside the SEQ field after separate and replace its number
    in_result = False
    for r in p._p.xpath("./w:r"):
        fld = r.find(qn("w:fldChar"))
        if fld is not None:
            ftype = fld.get(qn("w:fldCharType"))
            if ftype == "separate":
                in_result = True
            elif ftype == "end":
                in_result = False
        elif in_result:
            t = r.find(qn("w:t"))
            if t is not None and t.text and t.text.strip().isdigit():
                t.text = str(new_num)
                return True
    return False


def main():
    if not os.path.exists(BACKUP_DOCX):
        shutil.copy2(DESKTOP_DOCX, BACKUP_DOCX)
        print(f"Backed up original file to {BACKUP_DOCX}")

    doc = docx.Document(BACKUP_DOCX)

    # Find paragraph '3.2 หน้าจอเว็บแอปพลิเคขัน' and '4.\tสถาปัตยกรรมระบบ (System Architecture)' in body (skip TOC)
    idx_32 = None
    idx_40 = None
    for i, p in enumerate(doc.paragraphs):
        if p.style.name.lower().startswith("toc"):
            continue
        if "3.2 หน้าจอเว็บแอปพลิเค" in p.text:
            idx_32 = i
        elif (p.text.startswith("4.\tสถาปัตยกรรมระบบ") or "สถาปัตยกรรมระบบ (System Architecture)" in p.text) and p.style.name == "Heading 1":
            if idx_32 is not None and i > idx_32:
                idx_40 = i
                break

    if idx_32 is None or idx_40 is None:
        raise RuntimeError(f"Could not locate section boundaries: idx_32={idx_32}, idx_40={idx_40}")

    print(f"Found 3.2 at paragraph {idx_32}, section 4 at paragraph {idx_40}")

    # Convert any empty Caption paragraphs right before idx_32 (e.g. 491..494) to Normal
    for i in range(max(0, idx_32 - 5), idx_32):
        if doc.paragraphs[i].style.name == "Caption" and not doc.paragraphs[i].text.strip():
            doc.paragraphs[i].style = doc.styles["Normal"]

    # Style 3.2 heading to match 3.1 Heading 2
    p_32 = doc.paragraphs[idx_32]
    p_32.style = doc.styles["Heading 2"]
    p_32.alignment = None
    for r in p_32.runs:
        set_thai_font(r, 16, bold=True)

    # Remove placeholder paragraphs between idx_32 and idx_40
    for i in range(idx_40 - 1, idx_32, -1):
        p_del = doc.paragraphs[i]
        p_del._element.getparent().remove(p_del._element)

    # Re-locate section 4 paragraph reference
    ref_p = None
    for p in doc.paragraphs:
        if p.text.startswith("4.\tสถาปัตยกรรมระบบ") or "สถาปัตยกรรมระบบ (System Architecture)" in p.text:
            if p.style.name == "Heading 1":
                ref_p = p
                break

    if ref_p is None:
        raise RuntimeError("Could not find Heading 1 for section 4 after cleanup")

    # Insert all 12 subsections before ref_p
    for item in ADMIN_SCREENS:
        # 1. Subheading (e.g., 3.2.1 หน้าเข้าสู่ระบบ)
        p_sub = insert_paragraph_before(ref_p, "Normal")
        p_sub.paragraph_format.first_line_indent = Twips(990)
        r_sub = p_sub.add_run(f"{item['sub_no']} {item['title']}")
        set_thai_font(r_sub, 16, bold=False)

        # 2. Description (starting with tab + หน้าที่ผู้ดูแลระบบ...)
        p_desc = insert_paragraph_before(ref_p, "Normal")
        p_desc.paragraph_format.first_line_indent = Twips(990)
        r_tab = p_desc.add_run("\t")
        set_thai_font(r_tab, 16, bold=False)
        r_desc = p_desc.add_run(item["desc"])
        set_thai_font(r_desc, 16, bold=False)

        # 3. Screenshot image
        img_path = os.path.join(SCREENSHOT_DIR, item["img"])
        p_img = insert_paragraph_before(ref_p, "Normal")
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.keep_with_next = True
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Inches(5.8))

        # 4. Caption (ภาพที่ XX หน้า...)
        p_cap = insert_paragraph_before(ref_p, "Caption")
        add_seq_caption(p_cap, item["fig_no"], item["title"])

        # 5. Blank spacer line
        p_blank = insert_paragraph_before(ref_p, "Normal")
        r_blank = p_blank.add_run("")
        set_thai_font(r_blank, 16, bold=False)

    # Update subsequent figure numbers in section 4 and Appendix A (47 -> 58, 48 -> 59)
    for p in doc.paragraphs:
        if "สถาปัตยกรรมระบบ (Deployment Diagram) ของระบบบริหารกระบวนการผลิตสื่อ" in p.text and p.style.name == "Caption":
            update_seq_number_in_paragraph(p, 58)
            print("Updated Deployment Diagram caption to ภาพที่ 58")
        elif "แผนภาพกระบวนการทางธุรกิจในการบริหารจัดการผลิตสื่อ (BPMN Diagram)" in p.text and p.style.name == "Caption":
            update_seq_number_in_paragraph(p, 59)
            print("Updated BPMN Diagram caption to ภาพที่ 59")

    # Save to repo academic directory
    os.makedirs(os.path.dirname(REPO_DOCX), exist_ok=True)
    doc.save(REPO_DOCX)
    print(f"Saved updated document to {REPO_DOCX}")

    # Save to Desktop (or fallback if open in Word)
    try:
        doc.save(DESKTOP_DOCX)
        print(f"Saved updated document directly to {DESKTOP_DOCX}")
    except PermissionError:
        doc.save(FALLBACK_DOCX)
        print(f"Desktop file is open in Word; saved to {FALLBACK_DOCX}")


if __name__ == "__main__":
    main()
