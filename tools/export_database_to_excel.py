#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export_database_to_excel.py
Reads the exported JSON dumps of MongoDB content_management database
and builds a clean, readable Excel workbook with one sheet per collection:
  - Users
  - Teams
  - Contents
  - Tasks
  - Ideas
  - TeamActivities
Output: docs/academic/Draftly-Database-Export.xlsx
"""

import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DUMP_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "master", "backend", "dump", "content_management"))
OUTPUT_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "docs", "academic", "Draftly-Database-Export.xlsx"))

def build_excel_export():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    PRIMARY_COLOR = "1E293B"
    PRIMARY_TEXT = "FFFFFF"
    BORDER_COLOR = "CBD5E1"
    ZEBRA_FILL = "F8FAFC"
    PASS_FILL = "DCFCE7"
    PASS_TEXT = "15803D"

    font_header = Font(name="Tahoma", size=9, bold=True, color=PRIMARY_TEXT)
    font_data = Font(name="Tahoma", size=9, color="1E293B")
    font_data_bold = Font(name="Tahoma", size=9, bold=True, color="1E293B")
    fill_header = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    fill_zebra = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_side = Side(border_style="thin", color=BORDER_COLOR)
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    collections = [
        ("Users", "users.json", ["_id", "username", "email", "firstName", "lastName", "role", "position", "status", "workingStatus", "isOnline", "teamId", "createdAt"]),
        ("Teams", "teams.json", ["_id", "name", "code", "joinCode", "leader", "status", "description", "members", "createdAt"]),
        ("Contents", "contents.json", ["_id", "title", "platform", "category", "status", "teamId", "creatorId", "targetPublishDate", "publishedUrl", "createdAt"]),
        ("Tasks", "tasks.json", ["_id", "title", "type", "status", "contentId", "teamId", "assignedTo", "deliverableUrl", "startedAt", "submittedAt", "createdAt"]),
        ("Ideas", "ideas.json", ["_id", "title", "platform", "category", "status", "teamId", "createdBy", "createdAt"]),
        ("TeamActivities", "teamactivities.json", ["_id", "teamId", "user", "action", "entityType", "entityId", "details", "createdAt"])
    ]

    for sheet_name, filename, fields in collections:
        filepath = os.path.join(DUMP_DIR, filename)
        if not os.path.exists(filepath):
            continue

        with open(filepath, "r", encoding="utf-8") as f:
            docs = json.load(f)

        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        # Headers
        for c_idx, field in enumerate(fields, start=1):
            cell = ws.cell(row=1, column=c_idx, value=field)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_cell

        # Rows
        for r_idx, doc in enumerate(docs, start=2):
            fill = fill_zebra if r_idx % 2 == 0 else fill_white
            for c_idx, field in enumerate(fields, start=1):
                val = doc.get(field, "")
                if isinstance(val, (dict, list)):
                    val = json.dumps(val, ensure_ascii=False)
                elif val is None:
                    val = ""
                else:
                    val = str(val)

                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = font_data_bold if c_idx == 1 else font_data
                cell.fill = fill
                cell.border = border_cell
                cell.alignment = align_center if c_idx in (1, 3, 6, 8, 9, 10) else align_left

        # Widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                lines = val.split("\n")
                for line in lines:
                    line_len = len(line)
                    if line_len > max_len:
                        max_len = line_len
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Excel database export created successfully at: {OUTPUT_PATH}")

if __name__ == "__main__":
    build_excel_export()
